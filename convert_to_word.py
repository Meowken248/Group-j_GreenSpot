import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_inline_text(paragraph, text, default_font_name="Calibri", default_font_size=10.5, default_color=RGBColor(51, 65, 85)):
    # Pattern to match bold, italic, inline code, math, links
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\$[^$]+\$|\[[^\]]+\]\([^)]+\))')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = default_font_name
        run.font.size = Pt(default_font_size)
        run.font.color.rgb = default_color
        
        if part.startswith('**') and part.endswith('**'):
            run.text = part[2:-2]
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif part.startswith('*') and part.endswith('*'):
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run.text = f" {part[1:-1]} "
            run.font.name = "Consolas"
            run.font.size = Pt(default_font_size * 0.92)
            run.font.color.rgb = RGBColor(4, 120, 87)
        elif part.startswith('$') and part.endswith('$'):
            run.text = part[1:-1]
            run.font.name = "Cambria Math" if "Math" in part else "Consolas"
            run.font.size = Pt(default_font_size)
            run.font.color.rgb = RGBColor(16, 185, 129)
            run.italic = True
        elif part.startswith('[') and ']' in part and '(' in part and part.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if m:
                run.text = m.group(1)
                run.font.color.rgb = RGBColor(5, 150, 105)
                run.underline = True
            else:
                run.text = part
        else:
            run.text = part

def convert_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    doc = Document()
    
    # Page Setup: Standard A4 or Letter, 0.8 inch margins for high density & elegance
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Add Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("GREENSPOT SYSTEM | BẢNG TỔNG HỢP 48 CHỨC NĂNG DEMO GIAO DIỆN")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Hệ Thống Mạng Xã Hội Sinh Thái & Giám Sát Đô Thị Thông Minh GreenSpot — 2026")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)
        
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()
        
        # Check code block
        if stripped.startswith('```'):
            if in_code_block:
                # End code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                crun = cp.add_run("\n".join(code_lines))
                crun.font.name = "Consolas"
                crun.font.size = Pt(8.5)
                crun.font.color.rgb = RGBColor(30, 41, 59)
                
                # set light border
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(f'''
                    <w:tcBorders {nsdecls("w")}>
                        <w:left w:val="single" w:sz="12" w:space="0" w:color="059669"/>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                    </w:tcBorders>
                ''')
                tcPr.append(tcBorders)
                
                # Add tiny spacing after code block
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)
                
                in_code_block = False
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Empty line
        if not stripped:
            i += 1
            continue
            
        # Horizontal Rule
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBorder = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBorder)
            i += 1
            continue
            
        # Table detection
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            # Parse table
            raw_rows = []
            for tl in table_lines:
                # check if separator row like | :--- | :---: |
                cells = [c.strip() for c in tl.strip('|').split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells):
                    continue
                raw_rows.append(cells)
                
            if raw_rows:
                num_cols = max(len(r) for r in raw_rows)
                tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="4", val="single")
                
                # Style headers
                for row_idx, rdata in enumerate(raw_rows):
                    row = tbl.rows[row_idx]
                    is_header = (row_idx == 0)
                    for col_idx in range(num_cols):
                        cell = row.cells[col_idx]
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        set_cell_margins(cell, top=120 if is_header else 80, 
                                               bottom=120 if is_header else 80, 
                                               left=120, right=120)
                        
                        txt = rdata[col_idx] if col_idx < len(rdata) else ""
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(0)
                        
                        if is_header:
                            set_cell_background(cell, "064E3B") # Deep Forest Emerald
                            format_inline_text(cp, f"**{txt}**", default_font_size=9.5, default_color=RGBColor(255, 255, 255))
                            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg)
                            
                            # Alignment logic
                            if col_idx in (0, 3, 5) or len(txt) <= 10:
                                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                
                            format_inline_text(cp, txt, default_font_size=9, default_color=RGBColor(30, 41, 59))
                            
                # Add spacing after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(6)
            continue
            
        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            format_inline_text(p, stripped[2:], default_font_size=20, default_color=RGBColor(6, 78, 59))
            p.runs[0].bold = True
            i += 1
            continue
            
        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            format_inline_text(p, stripped[3:], default_font_size=14, default_color=RGBColor(5, 150, 105))
            for r in p.runs:
                r.bold = True
            i += 1
            continue
            
        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            format_inline_text(p, stripped[4:], default_font_size=12, default_color=RGBColor(15, 23, 42))
            for r in p.runs:
                r.bold = True
            i += 1
            continue
            
        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            format_inline_text(p, stripped[5:], default_font_size=11, default_color=RGBColor(4, 120, 87))
            for r in p.runs:
                r.bold = True
            i += 1
            continue
            
        # Math block $$ ... $$
        if stripped.startswith('$$') and stripped.endswith('$$'):
            formula = stripped[2:-2].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "F0FDF4")
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            crun = cp.add_run(formula)
            crun.font.name = "Cambria Math"
            crun.font.size = Pt(10)
            crun.font.color.rgb = RGBColor(6, 78, 59)
            crun.italic = True
            
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:left w:val="single" w:sz="8" w:space="0" w:color="10B981"/>
                    <w:top w:val="none"/>
                    <w:right w:val="none"/>
                    <w:bottom w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
            i += 1
            continue
            
        # Blockquote >
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            format_inline_text(p, stripped[2:], default_font_size=10, default_color=RGBColor(71, 85, 105))
            for r in p.runs:
                r.italic = True
            pBorder = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="8" w:color="10B981"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBorder)
            i += 1
            continue
            
        # Bullet list item
        if stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, stripped[2:], default_font_size=10, default_color=RGBColor(51, 65, 85))
            i += 1
            continue
            
        # Sub-bullet list item
        if line.startswith('    * ') or line.startswith('    - ') or line.startswith('  * ') or line.startswith('  - '):
            sub_text = re.sub(r'^\s*[\*\-]\s*', '', line)
            p = doc.add_paragraph(style='List Bullet 2' if 'List Bullet 2' in [s.name for s in doc.styles] else 'List Bullet')
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            format_inline_text(p, sub_text, default_font_size=9.5, default_color=RGBColor(71, 85, 105))
            i += 1
            continue
            
        # Numbered list item
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, m_num.group(2), default_font_size=10, default_color=RGBColor(51, 65, 85))
            i += 1
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        format_inline_text(p, stripped, default_font_size=10, default_color=RGBColor(51, 65, 85))
        i += 1
        
    doc.save(docx_path)
    print(f"Successfully converted {md_path} -> {docx_path}")

if __name__ == '__main__':
    md_file = r"d:\Group-j_GreenSpot\DANH_SACH_CHUC_NANG_DEMO_GIAO_DIEN.md"
    docx_file = r"d:\Group-j_GreenSpot\DANH_SACH_CHUC_NANG_DEMO_GIAO_DIEN.docx"
    convert_markdown_to_docx(md_file, docx_file)
