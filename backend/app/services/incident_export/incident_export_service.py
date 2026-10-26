import os
import uuid
import math
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch, cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.interface.incident_export.incident_export_interface import IIncidentExportRepository
from app.models.incident_export.exported_file import ExportedIncidentFile
from app.schemas.incident_export.schemas import (
    ExportCreatedResponse,
    ExportMetadataOptionsResponse,
    ExportStatusResponse,
    IncidentExportFilterParams,
    IncidentPreviewItem,
    IncidentPreviewResponse,
    TriggerExportRequest,
)

# Đăng ký font tiếng Việt chuẩn cho ReportLab
FONT_NAME = "Helvetica"
FONT_BOLD = "Helvetica-Bold"
try:
    if os.path.exists("C:/Windows/Fonts/arial.ttf"):
        pdfmetrics.registerFont(TTFont("ArialVN", "C:/Windows/Fonts/arial.ttf"))
        FONT_NAME = "ArialVN"
    if os.path.exists("C:/Windows/Fonts/arialbd.ttf"):
        pdfmetrics.registerFont(TTFont("ArialVN-Bold", "C:/Windows/Fonts/arialbd.ttf"))
        FONT_BOLD = "ArialVN-Bold"
    elif FONT_NAME == "ArialVN":
        FONT_BOLD = "ArialVN"
except Exception:
    pass

# Thư mục lưu trữ tệp xuất an toàn
EXPORT_STORAGE_DIR = Path(__file__).resolve().parent.parent.parent.parent / "storage" / "exports"
EXPORT_STORAGE_DIR.mkdir(parents=True, exist_ok=True)


def format_file_size(size_bytes: int) -> str:
    """Định dạng dung lượng tệp thân thiện với người dùng (KB, MB)"""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"


class IncidentExportService:
    """
    Tầng nghiệp vụ (Business Logic) phục vụ tính năng STT 47:
    Xuất danh sách sự cố ra Excel hoặc PDF theo bộ lọc.
    """

    def __init__(self, repository: IIncidentExportRepository):
        self.repository = repository

    def _resolve_quick_range(self, quick_range: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
        """Tính toán khoảng ngày Từ ngày - Đến ngày khi người dùng chọn nhanh"""
        if not quick_range:
            return None, None

        today = date.today()
        qr = quick_range.upper().strip()

        if qr == "TODAY":
            today_str = today.strftime("%Y-%m-%d")
            return today_str, today_str
        elif qr == "THIS_WEEK":
            start_week = today - timedelta(days=today.weekday())
            end_week = start_week + timedelta(days=6)
            return start_week.strftime("%Y-%m-%d"), end_week.strftime("%Y-%m-%d")
        elif qr == "THIS_MONTH":
            start_month = today.replace(day=1)
            # Tính ngày cuối tháng
            if start_month.month == 12:
                end_month = date(start_month.year + 1, 1, 1) - timedelta(days=1)
            else:
                end_month = date(start_month.year, start_month.month + 1, 1) - timedelta(days=1)
            return start_month.strftime("%Y-%m-%d"), end_month.strftime("%Y-%m-%d")
        elif qr == "THIS_QUARTER":
            quarter = (today.month - 1) // 3 + 1
            first_month_of_quarter = 3 * (quarter - 1) + 1
            start_quarter = date(today.year, first_month_of_quarter, 1)
            if quarter == 4:
                end_quarter = date(today.year, 12, 31)
            else:
                end_quarter = date(today.year, first_month_of_quarter + 3, 1) - timedelta(days=1)
            return start_quarter.strftime("%Y-%m-%d"), end_quarter.strftime("%Y-%m-%d")

        return None, None

    def validate_date_range(
        self, from_date: Optional[str], to_date: Optional[str]
    ) -> Tuple[bool, Optional[str]]:
        """Kiểm tra tính hợp lệ của khoảng ngày (Ngày bắt đầu <= Ngày kết thúc)"""
        if from_date and to_date:
            try:
                dt_from = datetime.strptime(from_date.strip(), "%Y-%m-%d")
                dt_to = datetime.strptime(to_date.strip(), "%Y-%m-%d")
                if dt_from > dt_to:
                    return False, "Khoảng ngày không hợp lệ dẫn đến không có dữ liệu. Vui lòng kiểm tra lại."
            except ValueError:
                return False, "Định dạng ngày không hợp lệ (yêu cầu YYYY-MM-DD)."
        return True, None

    async def preview_incidents(
        self, params: IncidentExportFilterParams
    ) -> IncidentPreviewResponse:
        """
        Khối 'DANH SÁCH' (Màn 1/4):
        Xử lý bộ lọc, phân trang và kiểm tra các trạng thái biên.
        """
        # Nếu có chọn nhanh thì ghi đè khoảng ngày
        if params.quick_range:
            q_from, q_to = self._resolve_quick_range(params.quick_range)
            if q_from and q_to:
                params.from_date = q_from
                params.to_date = q_to

        # Kiểm tra tính hợp lệ của ngày
        is_valid_date, err_msg = self.validate_date_range(params.from_date, params.to_date)
        if not is_valid_date:
            return IncidentPreviewResponse(
                success=True,
                message=err_msg or "Khoảng ngày không hợp lệ dẫn đến không có dữ liệu. Vui lòng kiểm tra lại.",
                total_records=0,
                page=params.page,
                page_size=params.page_size,
                total_pages=0,
                items=[],
                is_empty=True,
                date_range_invalid=True,
                error_detail=err_msg,
            )

        # Lấy dữ liệu xem trước phân trang
        items_raw, total_count = await self.repository.get_preview_incidents(params)
        total_pages = math.ceil(total_count / params.page_size) if total_count > 0 else 0

        items = [IncidentPreviewItem(**item) for item in items_raw]

        if total_count == 0:
            return IncidentPreviewResponse(
                success=True,
                message="Không có dữ liệu sự cố nào phù hợp để xuất báo cáo.",
                total_records=0,
                page=params.page,
                page_size=params.page_size,
                total_pages=0,
                items=[],
                is_empty=True,
                date_range_invalid=False,
            )

        return IncidentPreviewResponse(
            success=True,
            message=f"Tìm thấy {total_count} bản ghi sự cố phù hợp với bộ lọc.",
            total_records=total_count,
            page=params.page,
            page_size=params.page_size,
            total_pages=total_pages,
            items=items,
            is_empty=False,
            date_range_invalid=False,
        )

    async def generate_export(
        self, request: TriggerExportRequest, user_id: Optional[uuid.UUID] = None
    ) -> ExportCreatedResponse:
        """
        Màn 2/4 & Màn 3/4:
        Tạo tệp Excel (.xlsx) hoặc PDF (.pdf) theo các tiêu chí bộ lọc đã áp dụng.
        Áp dụng safeguard 15 phút thời hạn tải tệp.
        """
        # Nếu chọn nhanh thì giải quyết khoảng ngày
        if request.quick_range:
            q_from, q_to = self._resolve_quick_range(request.quick_range)
            if q_from and q_to:
                request.from_date = q_from
                request.to_date = q_to

        # Kiểm tra ngày hợp lệ
        is_valid_date, err_msg = self.validate_date_range(request.from_date, request.to_date)
        if not is_valid_date:
            raise ValueError(err_msg or "Khoảng ngày không hợp lệ dẫn đến không có dữ liệu. Vui lòng kiểm tra lại.")

        filter_params = IncidentExportFilterParams(
            district=request.district,
            category=request.category,
            status=request.status,
            from_date=request.from_date,
            to_date=request.to_date,
            page=1,
            page_size=10000,
        )

        incidents = await self.repository.get_all_matching_incidents(filter_params)
        if not incidents:
            raise ValueError("Không có dữ liệu sự cố nào phù hợp để xuất báo cáo.")

        export_id = uuid.uuid4()
        now_dt = datetime.now()
        timestamp_str = now_dt.strftime("%Y%m%d_%H%M%S")
        export_format = request.export_format.upper().strip()

        if export_format == "PDF":
            file_name = f"BaoCao_SuCo_{timestamp_str}.pdf"
            file_path = EXPORT_STORAGE_DIR / f"{export_id}_{file_name}"
            mime_type = "application/pdf"
            self._create_pdf_report(str(file_path), incidents, request, now_dt)
        else:
            file_name = f"BaoCao_SuCo_{timestamp_str}.xlsx"
            file_path = EXPORT_STORAGE_DIR / f"{export_id}_{file_name}"
            mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            self._create_excel_report(str(file_path), incidents, request, now_dt)

        file_size_bytes = os.path.getsize(str(file_path))
        file_size_formatted = format_file_size(file_size_bytes)

        # Thời hạn hết hạn: 15 phút kể từ lúc tạo
        expires_at = datetime.now(timezone.utc) + timedelta(minutes=15)

        filters_summary = {
            "district": request.district or "Tất cả",
            "category": request.category or "Tất cả",
            "status": request.status or "Tất cả",
            "from_date": request.from_date or "Toàn thời gian",
            "to_date": request.to_date or "Toàn thời gian",
        }

        # Lưu bản ghi siêu dữ liệu tệp vào CSDL
        file_record = ExportedIncidentFile(
            export_id=export_id,
            user_id=user_id,
            file_name=file_name,
            file_type=export_format,
            mime_type=mime_type,
            file_size_bytes=file_size_bytes,
            file_size_formatted=file_size_formatted,
            file_path=str(file_path),
            filters_applied=filters_summary,
            total_records=len(incidents),
            expires_at=expires_at,
            is_downloaded=False,
            download_count=0,
        )
        await self.repository.save_exported_file(file_record)

        download_url = f"/api/v1/incidents/export/{export_id}/download"

        return ExportCreatedResponse(
            success=True,
            export_id=str(export_id),
            file_name=file_name,
            file_type=export_format,
            file_size_formatted=file_size_formatted,
            file_size_bytes=file_size_bytes,
            expires_at=expires_at,
            is_expired=False,
            download_url=download_url,
            message="ĐÃ TẠO TỆP. Tệp sẵn sàng tải về.",
        )

    def _create_excel_report(
        self,
        output_path: str,
        incidents: List[Dict[str, Any]],
        request: TriggerExportRequest,
        now_dt: datetime,
    ) -> None:
        """Tạo bảng tính Excel (.xlsx) với phong cách thiết kế sang trọng, chuyên nghiệp"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh Sách Sự Cố"

        # Đặt font & màu chủ đạo
        BRAND_COLOR = "1E5E3A"      # Xanh lục GreenSpot
        HEADER_TEXT_COLOR = "FFFFFF"
        LIGHT_BG = "F4FBF7"
        BORDER_GRAY = "D1D5DB"

        title_font = Font(name="Arial", size=15, bold=True, color="1E5E3A")
        sub_font = Font(name="Arial", size=10, italic=True, color="4B5563")
        meta_label_font = Font(name="Arial", size=10, bold=True, color="1F2937")
        meta_val_font = Font(name="Arial", size=10, color="374151")
        tbl_header_font = Font(name="Arial", size=10, bold=True, color=HEADER_TEXT_COLOR)
        cell_font = Font(name="Arial", size=9)

        thin_side = Side(style="thin", color=BORDER_GRAY)
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        fill_header = PatternFill(start_color=BRAND_COLOR, end_color=BRAND_COLOR, fill_type="solid")
        fill_zebra = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type="solid")

        # 1. Khối Tiêu Đề
        ws.merge_cells("A1:G1")
        ws["A1"] = "HỆ THỐNG GIÁM SÁT MÔI TRƯỜNG ĐÔ THỊ GREENSPOT - TP. HỒ CHÍ MINH"
        ws["A1"].font = Font(name="Arial", size=10, bold=True, color="047857")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:G2")
        ws["A2"] = "BÁO CÁO TRÍCH XUẤT DANH SÁCH SỰ CỐ MÔI TRƯỜNG"
        ws["A2"].font = title_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A3:G3")
        ws["A3"] = f"Thời điểm trích xuất: {now_dt.strftime('%d/%m/%Y %H:%M:%S')} | Số lượng sự cố: {len(incidents)} bản ghi"
        ws["A3"].font = sub_font
        ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

        # 2. Khối Thông Tin Bộ Lọc Đã Áp Dụng
        ws["A5"] = "Bộ lọc đã áp dụng:"
        ws["A5"].font = meta_label_font
        ws["B5"] = f"Quận/Huyện: {request.district or 'Tất cả'}"
        ws["B5"].font = meta_val_font
        ws["D5"] = f"Loại sự cố: {request.category or 'Tất cả'}"
        ws["D5"].font = meta_val_font
        ws["F5"] = f"Trạng thái: {request.status or 'Tất cả'}"
        ws["F5"].font = meta_val_font

        ws["B6"] = f"Từ ngày: {request.from_date or 'Toàn bộ'}"
        ws["B6"].font = meta_val_font
        ws["D6"] = f"Đến ngày: {request.to_date or 'Toàn bộ'}"
        ws["D6"].font = meta_val_font

        # 3. Tiêu Đề Bảng Dữ Liệu
        headers = ["STT", "Mã sự cố", "Loại sự cố", "Quận / Huyện", "Trạng thái", "Địa chỉ chi tiết", "Ngày ghi nhận"]
        start_row = 8
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            cell.font = tbl_header_font
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border

        # 4. Dữ Liệu Các Dòng Sự Cố
        for row_idx, item in enumerate(incidents, start=1):
            curr_row = start_row + row_idx
            is_zebra = (row_idx % 2 == 0)

            ws.cell(row=curr_row, column=1, value=row_idx).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=2, value=item.get("tracking_code", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=3, value=item.get("category_name", "")).alignment = Alignment(horizontal="left")
            ws.cell(row=curr_row, column=4, value=item.get("district_name", "")).alignment = Alignment(horizontal="left")
            ws.cell(row=curr_row, column=5, value=item.get("status", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=6, value=item.get("address_text", "")).alignment = Alignment(horizontal="left")
            ws.cell(row=curr_row, column=7, value=item.get("created_at", "")).alignment = Alignment(horizontal="center")

            for col_idx in range(1, 8):
                c = ws.cell(row=curr_row, column=col_idx)
                c.font = cell_font
                c.border = cell_border
                if is_zebra:
                    c.fill = fill_zebra

        # Tự động co giãn độ rộng cột
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < start_row:
                    continue
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        ws.column_dimensions["F"].width = 38  # Địa chỉ
        ws.column_dimensions["A"].width = 8   # STT
        ws.column_dimensions["B"].width = 16  # Mã sự cố

        wb.save(output_path)

    def _create_pdf_report(
        self,
        output_path: str,
        incidents: List[Dict[str, Any]],
        request: TriggerExportRequest,
        now_dt: datetime,
    ) -> None:
        """Tạo tệp PDF theo quy chuẩn văn bản hành chính Việt Nam"""
        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            leftMargin=1.5 * cm,
            rightMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
        )

        styles = getSampleStyleSheet()

        # Tùy biến styles sử dụng font Unicode tiếng Việt
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName=FONT_BOLD,
            fontSize=13,
            leading=17,
            alignment=1,  # Center
            textColor=colors.HexColor("#1E5E3A"),
            spaceAfter=6,
        )

        sub_header_style = ParagraphStyle(
            "SubHeader",
            parent=styles["Normal"],
            fontName=FONT_NAME,
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.HexColor("#4B5563"),
            spaceAfter=12,
        )

        national_header_left = ParagraphStyle(
            "NationalHeaderLeft",
            parent=styles["Normal"],
            fontName=FONT_BOLD,
            fontSize=9,
            leading=12,
            alignment=0,
            textColor=colors.HexColor("#1F2937"),
        )

        national_header_right = ParagraphStyle(
            "NationalHeaderRight",
            parent=styles["Normal"],
            fontName=FONT_BOLD,
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.HexColor("#1F2937"),
        )

        motto_style = ParagraphStyle(
            "MottoStyle",
            parent=styles["Normal"],
            fontName=FONT_BOLD,
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.HexColor("#1F2937"),
        )

        body_style = ParagraphStyle(
            "BodyVN",
            parent=styles["Normal"],
            fontName=FONT_NAME,
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#374151"),
        )

        body_bold = ParagraphStyle(
            "BodyVNBold",
            parent=styles["Normal"],
            fontName=FONT_BOLD,
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#1F2937"),
        )

        table_header_style = ParagraphStyle(
            "TableHeaderVN",
            parent=styles["Normal"],
            fontName=FONT_BOLD,
            fontSize=8.5,
            leading=11,
            alignment=1,
            textColor=colors.white,
        )

        table_cell_center = ParagraphStyle(
            "TableCellCenter",
            parent=styles["Normal"],
            fontName=FONT_NAME,
            fontSize=8,
            leading=10,
            alignment=1,
        )

        story = []

        # 1. Quốc hiệu - Tiêu ngữ & Đơn vị chủ quản
        header_data = [
            [
                Paragraph("ỦY BAN NHÂN DÂN TP. HỒ CHÍ MINH<br/><b>HỆ THỐNG GREENSPOT</b>", national_header_left),
                Paragraph("<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><u>Độc lập - Tự do - Hạnh phúc</u>", national_header_right),
            ]
        ]
        header_table = Table(header_data, colWidths=[8.5 * cm, 9.5 * cm])
        header_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 10))

        # 2. Tiêu đề Báo Cáo
        story.append(Paragraph("BÁO CÁO TRÍCH XUẤT DANH SÁCH SỰ CỐ MÔI TRƯỜNG ĐÔ THỊ", title_style))
        story.append(Paragraph(
            f"Thời điểm xuất: {now_dt.strftime('%d/%m/%Y %H:%M:%S')} — Tổng số: {len(incidents)} sự cố",
            sub_header_style
        ))

        # 3. Khối Thông Tin Bộ Lọc
        filter_data = [
            [
                Paragraph(f"<b>Quận/Huyện:</b> {request.district or 'Tất cả'}", body_style),
                Paragraph(f"<b>Loại sự cố:</b> {request.category or 'Tất cả'}", body_style),
            ],
            [
                Paragraph(f"<b>Trạng thái:</b> {request.status or 'Tất cả'}", body_style),
                Paragraph(f"<b>Khoảng thời gian:</b> {request.from_date or 'Toàn bộ'} đến {request.to_date or 'Toàn bộ'}", body_style),
            ]
        ]
        filter_table = Table(filter_data, colWidths=[9.0 * cm, 9.0 * cm])
        filter_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0FDF4")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#BBF7D0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(filter_table)
        story.append(Spacer(1, 12))

        # 4. Bảng Dữ Liệu Sự Cố (Mã | Loại | Quận | Trạng thái | Ngày ghi nhận)
        table_rows = [
            [
                Paragraph("<b>STT</b>", table_header_style),
                Paragraph("<b>Mã sự cố</b>", table_header_style),
                Paragraph("<b>Loại sự cố</b>", table_header_style),
                Paragraph("<b>Quận / Huyện</b>", table_header_style),
                Paragraph("<b>Trạng thái</b>", table_header_style),
                Paragraph("<b>Ngày ghi nhận</b>", table_header_style),
            ]
        ]

        for idx, inc in enumerate(incidents, start=1):
            table_rows.append([
                Paragraph(str(idx), table_cell_center),
                Paragraph(f"<b>{inc.get('tracking_code', '')}</b>", table_cell_center),
                Paragraph(inc.get("category_name", ""), body_style),
                Paragraph(inc.get("district_name", ""), body_style),
                Paragraph(inc.get("status", ""), table_cell_center),
                Paragraph(inc.get("created_at", ""), table_cell_center),
            ])

        col_widths = [1.2 * cm, 3.2 * cm, 4.2 * cm, 3.8 * cm, 3.0 * cm, 2.6 * cm]
        inc_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        inc_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E5E3A")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(inc_table)
        story.append(Spacer(1, 18))

        # 5. Khối Ký Tên Văn Bản Chuẩn
        sig_data = [
            [
                Paragraph("<b>NGƯỜI LẬP BÁO CÁO</b><br/><i>(Ký và ghi rõ họ tên)</i>", table_cell_center),
                Paragraph(f"<i>TP. Hồ Chí Minh, ngày {now_dt.day} tháng {now_dt.month} năm {now_dt.year}</i><br/><b>THỦ TRƯỞNG ĐƠN VỊ</b><br/><i>(Ký tên, đóng dấu)</i>", table_cell_center),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[9.0 * cm, 9.0 * cm])
        sig_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(sig_table)

        doc.build(story)

    async def get_export_status(self, export_id: uuid.UUID) -> ExportStatusResponse:
        """Kiểm tra trạng thái tệp và chính sách bảo mật hết hạn 15 phút (Màn 3/4)"""
        record = await self.repository.get_exported_file_by_id(export_id)
        if not record:
            raise ValueError("Không tìm thấy tệp xuất báo cáo.")

        is_exp = record.is_expired
        if is_exp:
            msg = "Liên kết tải tệp đã hết hạn do chính sách bảo mật. Vui lòng xuất lại tệp."
        else:
            msg = "ĐÃ TẠO TỆP. Tệp sẵn sàng tải về."

        return ExportStatusResponse(
            export_id=str(record.export_id),
            file_name=record.file_name,
            file_type=record.file_type,
            file_size_formatted=record.file_size_formatted,
            is_expired=is_exp,
            expires_at=record.expires_at,
            message=msg,
        )

    async def get_download_file(
        self, export_id: uuid.UUID
    ) -> Tuple[str, str, str]:
        """
        Xác thực tải tệp: Kiểm tra tồn tại và thời hạn 15 phút.
        Trả về: (file_path, file_name, mime_type)
        """
        record = await self.repository.get_exported_file_by_id(export_id)
        if not record:
            raise ValueError("Không tìm thấy tệp xuất báo cáo.")

        if record.is_expired:
            raise PermissionError("Liên kết tải tệp đã hết hạn do chính sách bảo mật. Vui lòng xuất lại tệp.")

        if not os.path.exists(record.file_path):
            raise FileNotFoundError("Tệp vật lý trên máy chủ không tồn tại hoặc đã được dọn dẹp.")

        # Cập nhật số lần tải
        await self.repository.mark_as_downloaded(export_id)

        return record.file_path, record.file_name, record.mime_type

    async def get_metadata_options(self) -> ExportMetadataOptionsResponse:
        """Cung cấp danh mục quận/huyện, danh mục sự cố cho bộ lọc Màn 1/4"""
        districts = await self.repository.get_available_districts()
        categories = await self.repository.get_available_categories()

        statuses = [
            {"value": "Tất cả", "label": "Tất cả"},
            {"value": "Chờ tiếp nhận", "label": "Chờ tiếp nhận"},
            {"value": "Đang xử lý", "label": "Đang xử lý"},
            {"value": "Đã hoàn tất", "label": "Đã hoàn tất"},
            {"value": "Đã đóng", "label": "Đã đóng"},
        ]

        quick_ranges = [
            {"value": "TODAY", "label": "Hôm nay"},
            {"value": "THIS_WEEK", "label": "Tuần này"},
            {"value": "THIS_MONTH", "label": "Tháng này"},
            {"value": "THIS_QUARTER", "label": "Quý này"},
        ]

        return ExportMetadataOptionsResponse(
            districts=["Tất cả"] + districts,
            categories=["Tất cả"] + categories,
            statuses=statuses,
            quick_ranges=quick_ranges,
        )
