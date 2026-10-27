"""
Unit & Integration Tests cho Admin System Activity Dashboard (STT 48).

File: BackEnd/tests/test_admin_dashboard.py
Kiểm thử toàn diện 3 tầng:
  1. Service Business Logic (% tăng trưởng, lấp đầy ngày thiếu, xuất Excel, xuất PDF)
  2. Router Transport Endpoints (GET /overview, GET /daily-table, POST /export)
  3. Xử lý ngoại lệ, validation lỗi nghiệp vụ
"""

from __future__ import annotations

import io
from datetime import date, timedelta
from typing import List, Tuple
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.admin_dashboard.router import router as admin_dashboard_router
from app.interface.admin_dashboard.dashboard_interface import IAdminDashboardRepository
from app.schemas.admin_dashboard.dashboard_schema import ExportRequest
from app.services.admin_dashboard.dashboard_service import AdminDashboardService


# ─────────────────────────────────────────────────────────────────────────────
# MOCK REPOSITORY
# ─────────────────────────────────────────────────────────────────────────────

class MockAdminDashboardRepository(IAdminDashboardRepository):
    """Giả lập Repository theo đúng hợp đồng IAdminDashboardRepository (DIP)."""

    def __init__(self):
        # Dữ liệu mẫu 3 ngày
        self.sample_rows: List[Tuple[date, int, int, int]] = [
            (date(2026, 10, 20), 10, 25, 2),
            (date(2026, 10, 21), 15, 30, 5),
            (date(2026, 10, 22), 20, 45, 3),
        ]

    async def get_daily_stats(
        self, db: AsyncSession, start_date: date, end_date: date
    ) -> List[Tuple[date, int, int, int]]:
        return [r for r in self.sample_rows if start_date <= r[0] <= end_date]

    async def get_period_totals(
        self, db: AsyncSession, start_date: date, end_date: date
    ) -> Tuple[int, int, int]:
        filtered = [r for r in self.sample_rows if start_date <= r[0] <= end_date]
        total_u = sum(r[1] for r in filtered)
        total_p = sum(r[2] for r in filtered)
        total_g = sum(r[3] for r in filtered)
        return total_u, total_p, total_g

    async def get_previous_period_totals(
        self, db: AsyncSession, start_date: date, end_date: date
    ) -> Tuple[int, int, int]:
        # Giả lập kỳ trước có số liệu để test % tăng trưởng
        return 30, 60, 5

    async def get_daily_table(
        self,
        db: AsyncSession,
        start_date: date,
        end_date: date,
        page: int,
        page_size: int,
    ) -> Tuple[List[Tuple[date, int, int, int]], int]:
        filtered = [r for r in self.sample_rows if start_date <= r[0] <= end_date]
        total_count = len(filtered)
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size
        return filtered[start_idx:end_idx], total_count


# ─────────────────────────────────────────────────────────────────────────────
# 1. TEST SERVICE LAYER (Business Logic)
# ─────────────────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_get_activity_dashboard_overview():
    """Kiểm tra logic tính KPI, tăng trưởng % và lấp đầy ngày cho biểu đồ đường."""
    repo = MockAdminDashboardRepository()
    service = AdminDashboardService(repository=repo)
    mock_db = AsyncMock(spec=AsyncSession)

    res = await service.get_activity_dashboard(
        db=mock_db,
        period_days=30,
        start_date=date(2026, 10, 20),
        end_date=date(2026, 10, 22),
    )

    # 1. KPI cards
    assert res.kpi.total_new_users == 45  # 10 + 15 + 20
    assert res.kpi.total_new_posts == 100  # 25 + 30 + 45
    assert res.kpi.total_new_groups == 10  # 2 + 5 + 3

    # % tăng trưởng so với kỳ trước (30, 60, 5):
    # users: (45 - 30)/30 * 100 = +50.0%
    assert res.kpi.users_growth_pct == 50.0
    # posts: (100 - 60)/60 * 100 = +66.67%
    assert res.kpi.posts_growth_pct == 66.67
    # groups: (10 - 5)/5 * 100 = +100.0%
    assert res.kpi.groups_growth_pct == 100.0

    # 2. Daily series
    assert len(res.daily_series) == 3
    assert res.daily_series[0].stat_date == date(2026, 10, 20)
    assert res.daily_series[0].new_users == 10


@pytest.mark.asyncio
async def test_get_activity_dashboard_growth_zero_previous():
    """Kiểm tra xử lý chia cho 0 an toàn khi kỳ trước không có dữ liệu."""
    repo = MockAdminDashboardRepository()
    # Mock kỳ trước trả về 0 cho tất cả
    repo.get_previous_period_totals = AsyncMock(return_value=(0, 0, 0))
    service = AdminDashboardService(repository=repo)
    mock_db = AsyncMock(spec=AsyncSession)

    res = await service.get_activity_dashboard(
        db=mock_db,
        start_date=date(2026, 10, 20),
        end_date=date(2026, 10, 22),
    )

    assert res.kpi.users_growth_pct == 100.0
    assert res.kpi.posts_growth_pct == 100.0
    assert res.kpi.groups_growth_pct == 100.0


@pytest.mark.asyncio
async def test_get_daily_table_service():
    """Kiểm tra bảng số liệu có phân trang và hàng tổng kết / trung bình."""
    repo = MockAdminDashboardRepository()
    service = AdminDashboardService(repository=repo)
    mock_db = AsyncMock(spec=AsyncSession)

    res = await service.get_daily_report(
        db=mock_db,
        from_date=date(2026, 10, 20),
        to_date=date(2026, 10, 22),
        page=1,
        page_size=10,
    )

    assert res.total_records == 3
    assert len(res.rows) == 3
    assert res.summary.total_users == 45
    assert res.summary.total_posts == 100
    assert res.summary.total_groups == 10
    assert res.summary.avg_users_per_day == 15.0  # 45 / 3


@pytest.mark.asyncio
async def test_get_daily_table_invalid_dates():
    """Kiểm tra validate khi ngày bắt đầu lớn hơn ngày kết thúc."""
    repo = MockAdminDashboardRepository()
    service = AdminDashboardService(repository=repo)
    mock_db = AsyncMock(spec=AsyncSession)

    with pytest.raises(ValueError) as exc:
        await service.get_daily_report(
            db=mock_db,
            from_date=date(2026, 10, 25),
            to_date=date(2026, 10, 20),
        )
    assert "Ngày bắt đầu không được lớn hơn ngày kết thúc" in str(exc.value)


@pytest.mark.asyncio
async def test_export_excel():
    """Kiểm tra xuất tệp Excel (.xlsx) hợp lệ và đầy đủ tiêu đề cột."""
    repo = MockAdminDashboardRepository()
    service = AdminDashboardService(repository=repo)
    mock_db = AsyncMock(spec=AsyncSession)

    request = ExportRequest(
        from_date=date(2026, 10, 20),
        to_date=date(2026, 10, 22),
        format="xlsx",
    )

    file_bytes, filename, content_type = await service.export_report(mock_db, request)

    assert filename.endswith(".xlsx")
    assert "spreadsheetml.sheet" in content_type
    assert len(file_bytes) > 0

    # Đọc lại bằng openpyxl để xác minh cấu trúc file Excel
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    assert ws.title == "Thống Kê Hoạt Động"
    # Dòng 1: Tiêu đề báo cáo
    assert "BÁO CÁO THỐNG KÊ HOẠT ĐỘNG HỆ THỐNG GREENSPOT" in str(ws["A1"].value)
    # Dòng 4: Header các cột
    assert ws["A4"].value == "Ngày"
    assert ws["B4"].value == "Người Dùng Mới"
    assert ws["C4"].value == "Bài Đăng Mới"
    assert ws["D4"].value == "Nhóm Mới"


@pytest.mark.asyncio
async def test_export_pdf():
    """Kiểm tra xuất tệp PDF (.pdf) hợp lệ."""
    repo = MockAdminDashboardRepository()
    service = AdminDashboardService(repository=repo)
    mock_db = AsyncMock(spec=AsyncSession)

    request = ExportRequest(
        from_date=date(2026, 10, 20),
        to_date=date(2026, 10, 22),
        format="pdf",
    )

    file_bytes, filename, content_type = await service.export_report(mock_db, request)

    assert filename.endswith(".pdf")
    assert content_type == "application/pdf"
    # PDF magic header bytes
    assert file_bytes.startswith(b"%PDF")


# ─────────────────────────────────────────────────────────────────────────────
# 2. TEST ROUTER LAYER (FastAPI Endpoints)
# ─────────────────────────────────────────────────────────────────────────────

def get_test_app() -> FastAPI:
    app = FastAPI()
    app.include_router(admin_dashboard_router, prefix="/api/v1")
    return app


def test_api_overview_endpoint(monkeypatch):
    """Kiểm tra endpoint GET /api/v1/admin-dashboard/overview."""
    app = get_test_app()
    client = TestClient(app)

    # Mock service singleton
    mock_service = AdminDashboardService(repository=MockAdminDashboardRepository())
    monkeypatch.setattr("app.api.v1.admin_dashboard.router._service", mock_service)

    # Mock database session dependency
    from app.database import get_db
    async def override_get_db():
        yield AsyncMock(spec=AsyncSession)

    app.dependency_overrides[get_db] = override_get_db

    response = client.get("/api/v1/admin-dashboard/overview?period_days=30")
    assert response.status_code == 200
    data = response.json()
    assert "kpi" in data
    assert "daily_series" in data
    assert data["period_days"] == 30


def test_api_daily_table_endpoint(monkeypatch):
    """Kiểm tra endpoint GET /api/v1/admin-dashboard/daily-table."""
    app = get_test_app()
    client = TestClient(app)

    mock_service = AdminDashboardService(repository=MockAdminDashboardRepository())
    monkeypatch.setattr("app.api.v1.admin_dashboard.router._service", mock_service)

    from app.database import get_db
    async def override_get_db():
        yield AsyncMock(spec=AsyncSession)

    app.dependency_overrides[get_db] = override_get_db

    response = client.get(
        "/api/v1/admin-dashboard/daily-table?from_date=2026-10-20&to_date=2026-10-22&page=1&page_size=15"
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_records"] == 3
    assert len(data["rows"]) == 3
    assert data["summary"]["total_users"] == 45


def test_api_export_endpoint_xlsx(monkeypatch):
    """Kiểm tra endpoint POST /api/v1/admin-dashboard/export trả về file Excel."""
    app = get_test_app()
    client = TestClient(app)

    mock_service = AdminDashboardService(repository=MockAdminDashboardRepository())
    monkeypatch.setattr("app.api.v1.admin_dashboard.router._service", mock_service)

    from app.database import get_db
    async def override_get_db():
        yield AsyncMock(spec=AsyncSession)

    app.dependency_overrides[get_db] = override_get_db

    payload = {
        "from_date": "2026-10-20",
        "to_date": "2026-10-22",
        "format": "xlsx",
    }
    response = client.post("/api/v1/admin-dashboard/export", json=payload)
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment;" in response.headers["content-disposition"]
    assert len(response.content) > 0


def test_api_export_endpoint_pdf(monkeypatch):
    """Kiểm tra endpoint POST /api/v1/admin-dashboard/export trả về file PDF."""
    app = get_test_app()
    client = TestClient(app)

    mock_service = AdminDashboardService(repository=MockAdminDashboardRepository())
    monkeypatch.setattr("app.api.v1.admin_dashboard.router._service", mock_service)

    from app.database import get_db
    async def override_get_db():
        yield AsyncMock(spec=AsyncSession)

    app.dependency_overrides[get_db] = override_get_db

    payload = {
        "from_date": "2026-10-20",
        "to_date": "2026-10-22",
        "format": "pdf",
    }
    response = client.post("/api/v1/admin-dashboard/export", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
