import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document("Báo Cáo (3).docx")

with open("scratch/features_25_27_extracted.txt", "w", encoding="utf-8") as out:
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if any(k in t.lower() for k in ['thử thách xanh', 'audit log', 'nhật ký kiểm toán', 'thư viện sống xanh', 'cẩm nang phân loại', 'chức năng 25', 'chức năng 26', 'chức năng 27', '25.', '26.', '27.']):
            out.write(f"P[{i}]: {t}\n")

    out.write("\n=== DETAILED PARAGRAPH RANGE DUMP ===\n")
    # Let's see what is around paragraphs 1350 to 1550 (as noted in summary)
    for i in range(1350, 1550):
        if i < len(doc.paragraphs):
            t = doc.paragraphs[i].text.strip()
            if t:
                out.write(f"P[{i}]: {t}\n")

    out.write("\n=== TABLES ===\n")
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_text = " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
            if any(k in row_text.lower() for k in ['25', '26', '27', 'audit log', 'thử thách', 'cẩm nang']):
                out.write(f"T[{t_idx}][{r_idx}]: {row_text}\n")

print("Done writing scratch/features_25_27_extracted.txt")
