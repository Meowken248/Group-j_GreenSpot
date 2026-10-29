# -*- coding: utf-8 -*-
import docx
import os
import re

doc = docx.Document(r"c:\wamp64\www\Group-j_GreenSpot\Báo Cáo.docx")

rels = doc.part.rels

with open(r"c:\wamp64\www\Group-j_GreenSpot\scratch\all_docx_images.txt", "w", encoding="utf-8") as out:
    for idx, p in enumerate(doc.paragraphs):
        r_ids = []
        for run in p.runs:
            if 'drawing' in run._r.xml:
                embeds = re.findall(r'r:embed="([^"]+)"', run._r.xml)
                r_ids.extend(embeds)
        if r_ids:
            # find surrounding text
            prev_text = doc.paragraphs[idx-1].text.strip() if idx > 0 else ""
            curr_text = p.text.strip()
            next_text = doc.paragraphs[idx+1].text.strip() if idx < len(doc.paragraphs)-1 else ""
            target_files = [rels[r].target_ref for r in r_ids if r in rels]
            out.write(f"P[{idx}] -> Imgs: {target_files}\n")
            out.write(f"   Prev: {prev_text}\n")
            out.write(f"   Curr: {curr_text}\n")
            out.write(f"   Next: {next_text}\n\n")

print("Done scanning all docx images!")
