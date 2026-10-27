# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"c:\wamp64\www\Group-j_GreenSpot\Báo Cáo.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

found_paras = []
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if any(k in text for k in ["Huỳnh Anh Tú", "24211TT0857", "Chức năng 13", "Chức năng 14", "Chức năng 24", "4.13", "4.14", "4.24"]):
        found_paras.append((idx, text[:120]))

for idx, t in found_paras[:30]:
    try:
        print(f"P[{idx}]: {t}")
    except Exception:
        print(f"P[{idx}]: [encode err]")

# Check tables as well
found_tables = []
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        row_text = " | ".join(c.text.strip() for c in row.cells)
        if any(k in row_text for k in ["Huỳnh Anh Tú", "24211TT0857", "13", "14", "24"]):
            if any(k in row_text for k in ["Tú", "24211TT0857"]):
                found_tables.append((t_idx, r_idx, row_text[:120]))

print("\n--- Tables with Tu ---")
for t_idx, r_idx, txt in found_tables[:20]:
    try:
        print(f"T[{t_idx}] R[{r_idx}]: {txt}")
    except Exception:
        pass
