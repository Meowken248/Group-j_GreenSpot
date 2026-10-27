# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"c:\wamp64\www\Group-j_GreenSpot\Báo Cáo.docx")

# Let's inspect the paragraphs around 310 to 550 and see what image relationships exist
with open(r"c:\wamp64\www\Group-j_GreenSpot\scratch\tu_images_map.txt", "w", encoding="utf-8") as out:
    for i in range(300, 560):
        if i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            t = p.text.strip()
            # check if paragraph has images (runs with drawings)
            img_ids = []
            for run in p.runs:
                if 'drawing' in run._r.xml:
                    # extract blip r:embed
                    import re
                    embeds = re.findall(r'r:embed="([^"]+)"', run._r.xml)
                    img_ids.extend(embeds)
            
            if t or img_ids:
                out.write(f"P[{i}]: {t} | Imgs: {img_ids}\n")

print("Done mapping tu images!")
